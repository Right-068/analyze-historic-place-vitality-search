"""Read-only raw reconstruction; original V1.4.10 ledger/template contracts."""
from __future__ import annotations
import csv,json,os,tempfile,re,zipfile,io
from collections import Counter
from pathlib import Path
from urllib.parse import urlparse
import ctrip
from common import *
from collect import Store

REVIEW_FIELDS=["任务运行编号","内部样本编号","所属来源页面编号","平台","景点ID","携程景点ID","平台景点名称",
 "评论ID","评论全文","字数","语言","是否平台译文","外文原文","平台译文","用户原文","总评分","分项评分","发布时间",
 "来源标记","出游类型","IP属地","有用数","回复数","图片数","视频数","商家回复","作者ID","评论详情页",
 "抓取时间","原始文件","原始文件SHA256","原始定位","来源参数"]

def fields(template):
    with (SKILL_DIR/"assets/templates"/template).open(encoding="utf-8-sig",newline="") as f: return next(csv.reader(f))
def compatible(row,header): return {k:row.get(k,"") for k in header}
def reconstruct(data,cfg):
    data=Path(data); task=read_json(data/"task.json"); state=read_json(data/"logs/运行状态.json")
    store=Store(data,cfg); sources=[]; searches=[]; reviews={}; occurrences=[]; evidence=[]; errors=[]
    historical=state.get("historical_matches",{})
    sites={m["selected"]["poiId"]:m["selected"] for m in [*historical.values(),*state["matches"].values()] if m["status"]=="matched"}
    sites.update(state.get('historical_platform_sites',{}))
    sites.update({r['poi_id']:r['candidate'] for r in state.get('entity_mappings',[]) if r['allow_collection']})
    bypoi={}
    for m in state["matches"].values():
        if m["status"]=="matched": bypoi.setdefault(m["selected"]["poiId"],[]).append(m["object"])
    run=task["task_run_id"]
    for key,rec in sorted(store.units.items()):
        for attempt in rec.get("attempts",[]):
            try: store.verified(attempt)
            except IntegrityError as e: errors.append(key+":archived_attempt:"+str(e))
        execution="X-"+short_hash(key)
        searches.append({"task_run_id":run,"search_id":execution,"query_id":rec["query_id"],"query":rec["query"],
                         "record_type":"execution","execution_id":execution,"retrieval_entry":rec["url"],"search_tool":"携程公开接口",
                         "started_at":rec["started_at"],"finished_at":rec.get("finished_at",""),"retrieved_at":rec.get("retrieved_at",""),
                         "status":rec["status"],"notes":rec.get("error",""),"target_platform_id":"ctrip","iteration_round":0})
        if rec["status"]!="success": continue
        try: raw=store.verified(rec)
        except IntegrityError as e: errors.append(key+":"+str(e)); continue
        if rec["kind"]=="search": continue
        if rec.get("body",{}).get("arg",{}).get("sourceType")!=cfg["ctrip"]["source_type"]: continue
        pid=rec["poi_id"]
        if pid not in sites: errors.append(key+":no reliable matched POI"); continue
        site=sites[pid]; items,total=ctrip.items_of(raw); sid="S-"+short_hash(key+rec["sha256"])
        levels={o["level"] for o in bypoi.get(pid,[])}
        level="area_direct" if levels=={"area_direct"} else "poi"
        # An API response is one captured source, not one independently browsed detail page.
        sources.append({"task_run_id":run,"source_id":sid,"place_name":site["name"],"entity_level":level,
                        "platform":"携程","domain":urlparse(rec["url"]).hostname,"source_category":"user_review",
                        "page_title":site["name"]+" / 公开评论接口响应","retrieved_at":rec["retrieved_at"],"url":rec["url"],
                        "query_id":rec["query_id"],"query_text":rec["query"],"access_status":"full","content_layer":"user_review",
                        "selection_mechanism":"platform_public_api","is_user_source":"true","used_for_scoring":"false","score_scope":"not_scored",
                        "dedup_group":rec["sha256"],"evidence_locator":rec["raw_file"],"execution_id":execution,"raw_response_sha256":rec["sha256"],
                        "notes":"Part04.03采集材料，未执行正式证据准入；请求参数见采集单元.json"})
        for i,item in enumerate(items):
            cid=str(item["commentId"]); rid="E-"+short_hash("ctrip:"+pid+":"+cid)
            parsed=ctrip.parse_item(item,site["sight_id"])
            row={"任务运行编号":run,"内部样本编号":rid,"所属来源页面编号":sid,"平台":"携程","景点ID":pid,"携程景点ID":site["sight_id"],
                 "平台景点名称":site["name"],**parsed,"抓取时间":rec["retrieved_at"],"原始文件":rec["raw_file"],
                 "原始文件SHA256":rec["sha256"],"原始定位":f"/result/items/{i}","来源参数":rec["body"]["arg"]["sourceType"]}
            occurrences.append({"evidence_id":rid,"source_id":sid,"raw_file":rec["raw_file"],"sha256":rec["sha256"],
                                "locator":row["原始定位"],"text_sha256":short_hash(parsed["评论全文"],64)})
            if (pid,cid) in reviews: continue
            reviews[pid,cid]=row
            evidence.append({"task_run_id":run,"evidence_id":rid,"source_id":sid,"platform":"携程","unit_type":"review",
                             "internal_author_id":parsed["作者ID"],"published_at":parsed["发布时间"],"entity_level":level,
                             "original_summary_text":parsed["评论全文"],"excerpt_or_summary":parsed["评论全文"],
                             "native_rating_value":parsed["总评分"],"native_rating_scale_min":1 if parsed["总评分"] is not None else "",
                             "native_rating_scale_max":5 if parsed["总评分"] is not None else "",
                             "used_for_scoring":"false","score_scope":"not_scored","dedup_group":rid,
                             "notes":"Part04.03原始评论；正式相关性、有效性、主维度及评分留待独立后续阶段"})
    relations=[]
    for row in reviews.values():
        for obj in bypoi.get(row["景点ID"],[]):
            for aid in obj["area_ids"]:
                relation=obj.get("area_relations",{}).get(aid,obj.get("spatial_status","as_recorded"))
                relations.append({"evidence_id":row["内部样本编号"],"poi_id":row["景点ID"],"object_id":obj["object_id"],"area_id":aid,"entity_level":obj["level"],
                                  "spatial_status":relation,"scope_adoption":"boundary_candidate_not_confirmed_inside" if relation.startswith("boundary_") else obj.get("scope_adoption","as_recorded")})
    entity_relations=[]
    entities={e['entity_id']:e for e in state.get('entity_model',{}).get('entities',[])}
    for row in reviews.values():
        for m in state.get('entity_mappings',[]):
            if m['poi_id']!=row['景点ID'] or not m['allow_collection']:continue
            ent=entities[m['entity_id']]
            for aid in ent['study_area_ids']:
                entity_relations.append({'evidence_id':row['内部样本编号'],'source_id':row['所属来源页面编号'],'poi_id':m['poi_id'],
                    'entity_id':ent['entity_id'],'area_id':aid,'mapping_id':m['mapping_id'],'relation_type':m['relation_type'],
                    'association_type':'direct' if m['relation_type']=='same_as' else m['relation_type'],
                    'relation_direction':m['relation_direction'],'basis':m['basis'],
                    'spatial_relation':ent['object'].get('area_relations',{}).get(aid,m['spatial_relation']),
                    'platform_spatial_relation':m['spatial_relation']})
    return {"reviews":list(reviews.values()),"sources":sources,"searches":searches,"evidence":evidence,"occurrences":occurrences,"relations":relations,"entity_relations":entity_relations,"errors":errors}

ZH={"matched":"已匹配","unmatched":"未匹配","unresolved":"待核实","ambiguous":"匹配歧义","conflict":"匹配冲突","completed":"完成","partial":"部分完成","failed":"失败","blocked":"访问受限",
    "running":"进行中","success":"成功","retryable_failure":"可恢复失败","inside":"范围内","outside":"范围外","boundary_uncertain":"边界待核实","id_conflict":"标识冲突",
    "invalid_coordinate":"无效坐标","coordinate_conflict":"坐标冲突","poi":"独立地点","area_direct":"研究区整体","internal_node":"附属节点","external_related":"外部关联",
    "uncertain":"待核实","A":"核心候选","B":"文化关联","C":"空间辅助","D":"不采集","E":"待核实",
    "invalid_platform_identity":"平台编号与链接无效或不一致","not_sight":"非景点对象","name_not_exact_or_alias":"名称或别名未可靠对应",
    "city_conflict_or_missing":"城市不一致或缺少城市信息","district_conflict":"行政区冲突","spatial_conflict":"空间位置冲突",
    "exact_alias_city_spatial":"名称或别名、城市及空间位置一致","exact_alias_city_address":"名称或别名、城市及明确地址一致",
    "insufficient_independent_location_evidence":"独立位置佐证不足","source_access_restricted":"来源访问受限",
    "missing_total_count":"接口总量缺失","repeated_page":"返回重复分页","platform_total_changed":"平台总量发生变化",
    "empty_before_known_end":"达到已知总量前出现空页","request_failed":"请求失败","page_limit":"达到分页上限"}
def zh(v):
    if isinstance(v,bool): return "是" if v else "否"
    if isinstance(v,str) and ':' in v:
        prefix,detail=v.split(':',1)
        if prefix in ('cultural_function','auxiliary_name','documented_business_culture','business_subject_not_cultural','cultural_attribute_and_type','source_parent_id'):
            return {'cultural_function':'自身文化功能','auxiliary_name':'附属设施名称','documented_business_culture':'自身文化业务属性','business_subject_not_cultural':'普通业务主体','cultural_attribute_and_type':'自身文化属性与类别','source_parent_id':'来源父级编号'}[prefix]+'：'+detail
    return ZH.get(v,v) if isinstance(v,str) else v
ZH.update({"interface_empty":"接口确定性空结果（不代表平台无评论）","cache_missing":"缺少离线缓存", "boundary_inside":"边界附近初步在内", "boundary_outside":"边界附近初步在外", "search_queries_incomplete":"部分查询未执行或缺少缓存"})
ZH.update({"address_number_conflict":"同街门牌范围冲突","empty_confirmation_incomplete_or_conflicting":"接口空结果未完成确认或存在矛盾"})
ZH.update({'research_area':'法定研究区','area_expression':'区域旅游表达','historic_place':'历史文化地点','cultural_facility':'文化设施','cultural_living_space':'文化生活场所','auxiliary_node':'辅助节点','uncertain_entity':'实体待核实',
 'priority':'优先检索','regular':'常规检索','retain_only':'仅保留','identified':'已识别','retained_not_searched':'保留且不检索','not_searched':'尚未检索','linked_collectable':'已关联可采集','no_candidate':'未发现候选','no_applicable_interface':'无适用评论接口',
 'same_as':'同一实体','part_of':'组成关系','represents_area':'区域旅游表达（不等同法定范围）','related_to':'其他关联','different_from':'不同实体','located_in':'范围内','near_boundary':'临近边界',
 'platform_to_entity':'平台对象指向研究实体','entity_to_platform':'研究实体指向平台对象','associated':'已建立非同一关系','cached':'复用存档','executed':'已执行','not_available':'未取得',
 'boundary_name':'边界材料名称','map_name':'地图名称','provided_alias':'输入资料别名','related_object_name':'关联对象名称','verified_name_discovery':'资料名称发现','normalized_discovery':'规范化名称发现','exploratory_area_expression':'区域表达探索','city_disambiguation':'城市限定发现','related_name_discovery':'关联名称发现',
 'native_parent_location_qualifier':'原生父级地点限定','native_name_suffix_exact_address_and_position':'名称短前缀、明确地址与位置共同佐证','area_name_role_and_verified_point_not_equal_boundary':'区域名称角色与点位佐证；不等同边界',
 'independently_identified_entity_point_inside_legal_scope':'独立确认实体点位在范围内','identified_boundary_entity_not_confirmed_inside':'独立确认边界对象，不确定在范围内','native_platform_parent_id':'平台原生父子编号','declared_name_identifies_distinct_entity':'资料名称指向不同实体','multiple_identity_candidates':'多个身份候选未确定','contradictory_native_identity_observations':'平台身份元数据相互冲突','no_applicable_comment_interface':'无适用评论接口','supported_by_record':'资料支持'})
ZH.update({'sight':'景点','hotel':'酒店','food':'餐饮','restaurant':'餐饮','core_typecode':'核心类别编码','independent_cultural_typecode':'独立文化类别','auxiliary_type':'辅助类别','input_class_A':'输入核心类别','input_class_B':'输入文化关联类别','as_recorded':'按原记录','unknown':'未知',
 'original_name':'原始名称','provided_alias':'输入资料别名','store_location_qualifier_removed':'剥离店铺位置限定','area_pedestrian_suffix_discovery':'步行街表达探索','area_base_discovery':'区域基础名探索','cultural_building_suffix':'文化建筑表达','area_expression_rule':'区域表达规则','city_qualifier_rule':'城市限定规则','authoritative_geometry_point_relation':'权威边界与点位关系','alias_names_distinct_source_entity':'别名指向不同来源实体'})

ENTITY_LINK_FIELDS=['evidence_id','source_id','poi_id','entity_id','area_id','mapping_id','relation_type','relation_direction','basis','spatial_relation','platform_spatial_relation','association_type']
def platform_relation(relation):
    return 'platform_part_of' if relation=='part_of' else relation
ZH.update({'direct':'直接对应','platform_part_of':'平台对象为研究实体组成节点',
           'ordinary_lodging_not_area':'普通住宿不是区域公众表达','non_place_commercial_service':'非地点商业服务',
           'missing_platform_name':'缺少平台名称'})
ZH.update({'L1':'一级：区域文化实体','L2':'二级：独立文化地点','L3':'三级：文化商业空间','L4':'四级：辅助空间节点',
 'HIGH':'高','MEDIUM':'中','LOW':'低','unrelated':'无关','pending':'等待模型判断','pending_host':'等待宿主判断','filtered':'空间或身份过滤','valid':'有效判断',
 'near_point':'点位附近','spatial_proposal':'通过空间候选筛选','ordinary_business_not_area':'普通业务主体不能代表研究区','auxiliary_not_collection_target':'辅助节点不独立采集'})
ZH.update({'filtered':'基础身份或普通业务过滤','spatial_proposal':'候选进入名称判断',
 'eligible':'允许采集','location_or_identity_blocked':'位置或身份门禁未通过','relation_not_accepted':'关系未达到采集条件','not_selected':'尚未采用'})
ZH.update({'context_proposal':'候选进入平台表达判断','not_regional_entity':'不是区域实体',
 'no_supported_internal_location':'缺少内部节点位置佐证','no_supported_spatial_context':'缺少关联空间背景',
 'related_collection_disabled':'本次不采集其他关联','area_point_not_equal_boundary':'平台点位不能证明同一法定区域',
 'model_decision_not_reliable':'模型判断未达采用条件'})
def entity_tables(state,result):
    """Presentation projections only; original seven templates are not altered."""
    entities=state.get('entity_model',{}).get('entities',[])
    byid={e['entity_id']:e for e in entities}
    def flat(v):return json.dumps(v,ensure_ascii=False,sort_keys=True) if isinstance(v,(list,dict)) else zh(v)
    def table(name,header,rows):return name,header,[[flat(v) for v in row] for row in rows]
    return [
      table('空间文化实体',['实体编号','实体名称','实体类型','来源对象编号','研究区域编号','文化依据','检索优先级','实体状态','平台可检索状态','实体层级','采集优先级','资料正式名称','已提供明确别名','平台表达检索假设','探索查询'],
            [[e['entity_id'],e['name'],e['entity_type'],e['source_object_ids'],e['study_area_ids'],e['cultural_basis'],e['retrieval_priority'],e['entity_status'],e['platform_status'],e.get('entity_level',''),e.get('collect_priority',''),e.get('official_name',e['name']),e.get('verified_alias',[]),e.get('platform_expression',[]),e.get('exploration_query',[])] for e in entities]),
      table('实体名称来源',['实体编号','名称','名称类型','来源编号','资料确认','仅作检索线索'],[[r[k] for k in ('entity_id','name','name_type','name_source','verified','search_only')] for e in entities for r in e['names']]),
      table('实体关系',['来源实体编号','目标实体编号','关系类型','空间关系','关系依据','记录状态'],[[r[k] for k in ('source_entity_id','target_entity_id','relation_type','spatial_relation','basis','status')] for r in state.get('entity_model',{}).get('relations',[])]),
      table('平台表达与实体关联',['关联编号','实体编号','平台地点编号','景点编号','平台名称','平台对象类型','平台表达关系','关系方向','程序空间关系','关系判断或过滤依据','程序允许采集','坐标解释','来源查询','关系判断置信度','判断状态','模型标识','提示词版本','判断请求编号','采集优先级','关系适用门禁','不采集原因','采集资格','采集执行状态'],
            [[r.get(k,'') for k in ('mapping_id','entity_id','poi_id','sight_id','platform_name','platform_type','relation_type','relation_direction','spatial_relation','basis','allow_collection','coordinate_interpretation','source_queries','confidence','decision_status','model_version','prompt_version','request_id','collect_priority','collection_gate','collection_block_reason','collection_status')]+[state.get('collections',{}).get(r['poi_id'],{}).get('status','not_selected') if r['allow_collection'] else 'not_selected'] for r in state.get('entity_mappings',[])]),
      table('平台实体关系判断',['研究实体编号','研究实体名称','平台实体编号','平台实体名称','平台表达关系','置信度','判断依据','采用的资料类别','模型标识','提示词版本','判断时间','判断状态'],
            [[r['entity_id'],byid[r['entity_id']]['name'],r['platform_expression_id'],r['platform_name'],r['relation_type'],r['confidence'],r['basis'],r.get('judgment_evidence',[]),r['model_version'],r['prompt_version'],r.get('decision_time',''),r['decision_status']] for r in state.get('entity_mappings',[])]),
      table('实体检索查询',['实体编号','查询词','查询来源','发现目的','资料名称','执行状态','查询编号','原始文件','原始哈希'],[[r[k] for k in ('entity_id','query','name_source','purpose','verified_name','execution_status','query_id','raw_file','sha256')] for r in state.get('entity_queries',[])]),
      table('评论实体区域关联',['评论编号','来源编号','评论原属平台地点编号','研究实体编号','研究区域编号','关联编号','平台表达关系','关系方向','关系判断依据','研究实体与区域空间关系','平台对象与研究实体空间关系','评论关联类型'],[[r[k] for k in ENTITY_LINK_FIELDS] for r in result['entity_relations']])]

def machine_tables(state,result):
    """Additional relation views only; never duplicate raw comments or ledger schema."""
    mappings=state.get('entity_mappings',[]);entities=state.get('entity_model',{}).get('entities',[])
    relationships=[{'entity_A':r['source_entity_id'],'entity_B':r['target_entity_id'],'relation':r['relation_type'],
        'spatial_relation':r['spatial_relation'],'reason':r['basis'],'status':r['status']}
        for r in state.get('entity_model',{}).get('relations',[])]
    expressions=[]
    for e in entities:
        for n in e['names']:
            expressions.append({'entity_id':e['entity_id'],'platform_name':n['name'],'expression_type':n['name_type'],
                'source':n['name_source'],'verified':'true' if n['verified'] else 'false'})
        for q in e.get('platform_expression',[]):
            expressions.append({'entity_id':e['entity_id'],'platform_name':q,'expression_type':'exploratory_area_expression',
                'source':'area_expression_rule','verified':'false'})
    for r in mappings:
        expressions.append({'entity_id':r['entity_id'],'platform_name':r['platform_name'],'expression_type':'observed_platform_name',
            'source':r['platform_expression_id'],'verified':'true'})
    expressions=list({tuple(r.values()):r for r in expressions}.values())
    platform=[{'research_entity':r['entity_id'],'ctrip_entity':r['platform_expression_id'],'relation':platform_relation(r['relation_type']),
        'collect_status':r['collection_status'],'collection_execution_status':state.get('collections',{}).get(r['poi_id'],{}).get('status','not_selected') if r['allow_collection'] else 'not_selected',
        'poi_id':r['poi_id'],'sight_id':r['sight_id'],'mapping_id':r['mapping_id'],
        'confidence':r['confidence'],'reason':r['basis'],'model':r['model_version'],'prompt_version':r['prompt_version'],
        'decision_status':r['decision_status'],'request_id':r['request_id']} for r in mappings]
    return [
        ('entity_relationship.csv',['entity_A','entity_B','relation','spatial_relation','reason','status'],relationships),
        ('platform_expression.csv',['entity_id','platform_name','expression_type','source','verified'],expressions),
        ('platform_mapping.csv',['research_entity','ctrip_entity','relation','collect_status','collection_execution_status','poi_id','sight_id','mapping_id','confidence','reason','model','prompt_version','decision_status','request_id'],platform),
        ('comment_relation.csv',ENTITY_LINK_FIELDS,result['entity_relations'])]
def split_text(s,limit=15000):
    parts=[]; chunk=[]; size=0
    for c in s:
        n=2 if ord(c)>65535 else 1
        if size+n>limit: parts.append("".join(chunk)); chunk=[]; size=0
        chunk.append(c); size+=n
    parts.append("".join(chunk)); return parts
def xml_unsupported(s):
    return any(not(c in "\t\n\r" or 0x20<=ord(c)<=0xD7FF or 0xE000<=ord(c)<=0xFFFD or 0x10000<=ord(c)<=0x10FFFF) for c in s)

def preserve_xml_cr(path):
    """ElementTree backends may emit literal CR, which XML readers normalize to LF.
    Escape only serialized text nodes, never decode user-supplied _x000D_ strings.
    """
    from openpyxl.xml.functions import LXML
    output=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(Path(path).read_bytes())) as src,zipfile.ZipFile(output,"w",compression=zipfile.ZIP_DEFLATED) as dst:
            for member in src.infolist():
                blob=src.read(member.filename)
                if member.filename.startswith("xl/worksheets/") and member.filename.endswith(".xml"):
                    if not LXML and os.linesep=="\r\n": blob=blob.replace(b"\r\n",b"\n")
                    blob=re.sub(rb"(<t(?:\s[^>]*)?>)(.*?)(</t>)",lambda m:m[1]+m[2].replace(b"\r",b"&#13;")+m[3],blob,flags=re.S)
                dst.writestr(member,blob)
    return output.getvalue()

def workbook(path,sheets):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(suffix=".xlsx",dir=path.parent); os.close(fd)
    overflow=[]
    try:
        wb=Workbook(write_only=True)
        def append(ws,values,header=False):
            cells=[]
            for v in values:
                cell=WriteOnlyCell(ws,value=v)
                if isinstance(v,str): cell.data_type="s"
                cell.font=Font(name="Carlito",size=10,bold=header,color="FFFFFF" if header else "000000")
                cell.alignment=Alignment(vertical="center" if header else "top",wrap_text=True)
                if header: cell.fill=PatternFill("solid",fgColor="17324D")
                cells.append(cell)
            ws.append(cells)
        try:
            names=set()
            for name,header,rows in sheets:
                for offset in range(0,max(1,len(rows)),1048575):
                    sheetname=name if offset==0 else name[:24]+"_"+str(offset//1048575+1)
                    if sheetname in names: raise ValueError("duplicate worksheet")
                    names.add(sheetname); ws=wb.create_sheet(sheetname); ws.sheet_view.showGridLines=False; ws.freeze_panes="A2"
                    ws.row_dimensions[1].height=34; ws.sheet_format.defaultRowHeight=25.5
                    for col in range(1,len(header)+1): ws.column_dimensions[get_column_letter(col)].width=18
                    append(ws,header,True)
                    subset=rows[offset:offset+1048575]
                    for r,values in enumerate(subset,1):
                        out=[]
                        for c,v in enumerate(values):
                            if v is None or v=="": out.append(None); continue
                            if isinstance(v,(dict,list)): v=json.dumps(v,ensure_ascii=False)
                            if isinstance(v,bool): v="是" if v else "否"
                            if isinstance(v,(int,float)): out.append(v)
                            else:
                                v=str(v)
                                if xml_unsupported(v):
                                    ref="XML-"+short_hash(sheetname+f":{r}:{c}")
                                    for n,chunk in enumerate(split_text(json.dumps(v,ensure_ascii=True)),1): overflow.append([ref,n,chunk])
                                    v=f"XML不可表示字符，JSON转义原文见「长文本续表」：{ref}；CSV保留全文"
                                elif len(v.encode("utf-16-le"))//2>32767:
                                    ref="LONG-"+short_hash(sheetname+f":{r}:{c}")
                                    for n,chunk in enumerate(split_text(v),1): overflow.append([ref,n,chunk])
                                    v=f"完整长文本见「长文本续表」：{ref}；CSV保留全文"
                                out.append(v)
                        append(ws,out)
                    ws.auto_filter.ref=f"A1:{get_column_letter(len(header))}{max(2,len(subset)+1)}"
            if overflow:
                for start in range(0,len(overflow),1048575):
                    ws=wb.create_sheet("长文本续表"+("" if not start else str(start//1048575+1)))
                    append(ws,["长文本编号","段序","原文分段"],True)
                    for i,row in enumerate(overflow[start:start+1048575],1):
                        append(ws,row)
            wb.save(tmp)
        finally: wb.close()
        atomic_bytes(path,preserve_xml_cr(tmp))
    finally:
        if os.path.exists(tmp): os.unlink(tmp)

def template_headers():
    import openpyxl
    w=openpyxl.load_workbook(SKILL_DIR/"assets/templates/网络检索与量化编码_模板.xlsx",read_only=True)
    heads={s.title:[c.value for c in s[1]] for s in w}; w.close(); return heads

def export(data,cfg):
    data=Path(data); result=reconstruct(data,cfg)
    # Preserve readable partial results; the independent verifier remains invalid.
    if result["errors"]: write_json(data/"verify/原始文件异常.json",{"errors":result["errors"]})
    state=read_json(data/"logs/运行状态.json"); task=read_json(data/"task.json")
    write_rows(data/'reviews/评论与实体关联.csv',result['entity_relations'],ENTITY_LINK_FIELDS)
    for name,header,rows in entity_tables(state,result):write_csv(data/f'entities/{name}.csv',[dict(zip(header,r)) for r in rows],header)
    for name,header,rows in machine_tables(state,result):write_csv(data/'entities'/name,rows,header)
    for name,key in (("source","sources"),("evidence","evidence"),("search-log","searches")):
        template=name+"-ledger-template.csv" if name!="search-log" else "search-log-template.csv"
        header=fields(template); write_csv(data/f"ledgers/{name+'-ledger' if name!='search-log' else name}.csv",[compatible(r,header) for r in result[key]],header)
    for filename,key in (("原始评论","reviews"),("评论与地点关联","relations"),("原始评论出现记录","occurrences")):
        header=REVIEW_FIELDS if key=="reviews" else ["evidence_id","poi_id","object_id","area_id","entity_level","spatial_status","scope_adoption"] if key=="relations" else ["evidence_id","source_id","raw_file","sha256","locator","text_sha256"]
        write_rows(data/f"reviews/{filename}.csv",result[key],header)
    matches=[]; candidates=[]
    for oid,m in state["matches"].items():
        s=m.get("selected") or {}
        matches.append({"对象编号":oid,"对象名称":m["object"]["name"],"对象层级":zh(m["object"]["level"]),"匹配状态":zh(m["status"]),
                        "平台地点编号":s.get("poiId",""),"平台名称":s.get("name",""),"判定依据":zh(s.get("reason",m.get("reason",m["status"]))),
                        "关联区域":";".join(m["object"]["area_ids"])})
        for c in m.get("candidates",[]): candidates.append({"object_id":oid,**c})
    write_rows(data/"matching/携程POI匹配结果.csv",matches)
    write_rows(data/"matching/携程候选清单.csv",[{k:json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v for k,v in r.items()} for r in candidates])
    write_rows(data/"matching/未匹配及歧义地点.csv",[r for r in matches if r["匹配状态"]!="已匹配"])
    review_site={r["内部样本编号"]:r["景点ID"] for r in result["reviews"]}
    raw_counts=Counter(review_site.get(r["evidence_id"]) for r in result["occurrences"])
    collections=[{"平台地点编号":r["poi_id"],"平台地点名称":r["name"],"状态":zh(r["status"]),"接口报告总数":r.get("interface_reported_total"),
                  "页面显示总数":r.get("platform_displayed_total"),"去重已取得":sum(x["景点ID"]==r["poi_id"] for x in result["reviews"]),
                  "原始出现量":raw_counts[r["poi_id"]],
                  "展示与接口差额":r.get("platform_displayed_total")-r.get("interface_reported_total") if isinstance(r.get("platform_displayed_total"),(int,float)) and isinstance(r.get("interface_reported_total"),(int,float)) else None,
                  "无法取得数":r.get("unavailable_count"),"口径说明":"仅判断指定接口范围；展示差额不能直接认定为漏采", "限制说明":";".join(zh(p) for p in r.get("problems",[]))} for r in state["collections"].values()]
    heads=template_headers(); content={k:[] for k in heads}
    for r in result["reviews"]:
        mapped={"内部样本编号":r["内部样本编号"],"所属来源页面编号":r["所属来源页面编号"],"页面标题":r["平台景点名称"],"发布时间/评论时间":r["发布时间"],
                "内容类型":"用户评论","摘要原文":r["评论全文"],"必要摘录/匿名化摘要":r["评论全文"],"来源链接":r["评论详情页"],"来源类型":"用户评价",
                "内容层级":"用户评论","是否纳入量化评分":"否","原生评分值":r["总评分"],"原生量表下界":1,"原生量表上界":5,
                "任务运行编号":r["任务运行编号"],"备注":"仅采集；完整原文、译文、来源参数及文件哈希见原始评论表"}
        content["样本明细"].append([mapped.get(h,"") for h in heads["样本明细"]])
    for r in result["sources"]:
        mapped=dict(zip(heads["来源页面"],[r["source_id"],r["place_name"],zh(r["entity_level"]),"携程",r["domain"],"用户评价",r["page_title"],"",r["retrieved_at"],r["url"],r["query_id"],r["query_text"],"","已读取","用户评论","公开接口","","是","","","否","本阶段不评分",r["dedup_group"],"",r["evidence_locator"],r["notes"],r["task_run_id"]]))
        content["来源页面"].append([mapped.get(h,"") for h in heads["来源页面"]])
    for r in result["searches"]:
        m={"任务运行编号":r["task_run_id"],"检索记录编号":r["search_id"],"检索日期":r["retrieved_at"],"检索工具或入口":r["retrieval_entry"],"实际检索词":r["query"],
           "检索状态":zh(r["status"]),"备注":r["notes"],"记录类型":"实际执行","执行编号":r["execution_id"],"查询编号":r["query_id"],"执行开始时间":r["started_at"],"执行结束时间":r["finished_at"]}
        content["检索日志"].append([m.get(h,"") for h in heads["检索日志"]])
    content["数据限制说明"]=[
        ["阶段边界","仅空间解析、携程匹配和原始采集；不作正式证据资格或评分","评价字段留空","不得当作正式研究报告"],
        ["访问边界","受限来源立即停止，不切换路径规避","保存真实状态","不宣称全量采集"],
        ["坐标精度","近似坐标及平台未知坐标基准","边界附近与匹配冲突单列","空间归属存在不确定性"],
        ["隐私","结构化作者匿名；原件可能含昵称","原始目录限制分享","仅获授权研究使用"],
        ["工作表","沿用原版七表名称与表头；主题和评分表仅保留结构","本阶段不生成评价数值","机器层未执行指标留空"]]
    search_count=sum(r["kind"]=="search" for r in Store(data,cfg).units.values())
    machine={"task_run_id":task["task_run_id"],"source_rows":len(result["sources"]),"evidence_rows":len(result["reviews"]),"search_query_rows":search_count,"complete_comment_claim":"仅反映已取得数据"}
    content["机器可读汇总"]=[[machine.get(h,"") for h in heads["机器可读汇总"]]]
    sheets=[(name,header,content[name]) for name,header in heads.items()]
    for name,rows in (("原始评论",result["reviews"]),("地点匹配",matches),("采集结果",collections)):
        h=REVIEW_FIELDS if name=="原始评论" else list(rows[0]) if rows else ["状态"]; sheets.append((name,h,[[r.get(k,"") for k in h] for r in rows]))
    summary=state.get("spatial_summary",{})
    sheets.append(("空间汇总",["项目","数量"],[[{"districts":"研究街区","poi_input_rows":"输入POI行","unique_poi":"去重POI","spatial_links":"空间关联","search_objects":"检索对象"}.get(k,k),v] for k,v in summary.items() if k!="categories"]))
    sheets.append(("字段说明",["项目","说明"],[["模板","前七表采用原版模板表名与表头；后续评分字段只预留"],["原始响应","以内容SHA256命名，采集单元绑定请求参数、状态、时间"],["评论唯一性","来源平台＋平台地点ID＋评论ID；重复出现另表保留"],["长文本","超过单元格容量时分段续表，CSV始终保留全文"],["平台变化","同ID不同时间内容变化保留不同原始出现，不判为本地原件损坏"]]))
    sheets.extend(entity_tables(state,result))
    workbook(data/"网络检索与量化编码_Part04.03.xlsx",sheets)
    write_json(data/"verify/导出清单.json",{str(p.relative_to(data)).replace("\\","/"):digest_file(p) for folder in ("reviews","ledgers","matching","entities") for p in (data/folder).glob("*.csv")})
    return {"status":"exported","reviews":len(result["reviews"]),"sources":len(result["sources"])}

def verify(data,cfg):
    data=Path(data); expected=reconstruct(data,cfg); errors=list(expected["errors"]); w=None
    def compare(path,rows,header=None):
        p=data/path
        if not p.is_file(): errors.append("missing "+path); return
        actual=read_csv(p)
        ref=[{k:"" if v is None else str(v) for k,v in (compatible(r,header) if header else r).items()} for r in rows]
        if actual!=ref: errors.append("table differs from raw: "+path)
    compare("reviews/原始评论.csv",expected["reviews"])
    compare("reviews/评论与地点关联.csv",expected["relations"])
    compare("reviews/原始评论出现记录.csv",expected["occurrences"])
    compare('reviews/评论与实体关联.csv',expected['entity_relations'])
    state=read_json(data/'logs/运行状态.json')
    tables=entity_tables(state,expected)
    for name,header,rows in tables:compare(f'entities/{name}.csv',[dict(zip(header,r)) for r in rows])
    for name,header,rows in machine_tables(state,expected):compare('entities/'+name,rows,header)
    if 'entity_model' in state:
        import copy,entity_relation,llm_entity_match
        profile=read_json(data/'matching/坐标解释校准.json')
        try:
            relation_cfg={**cfg,'matching':{**cfg['matching'],'calibration':profile}}
            model=copy.deepcopy(state['entity_model'])
            proposed,_,requests=entity_relation.proposals(model,Store(data,cfg).cached_candidates(),relation_cfg)
            records=llm_entity_match.resolve(data,requests,relation_cfg,read_only=True)
            mappings=entity_relation.apply(model,proposed,records,relation_cfg)
            if mappings!=state.get('entity_mappings'):errors.append('entity mappings differ from archived search recomputation')
        except (IntegrityError,InvalidResponse,ValueError,KeyError) as e:errors.append('entity association:'+str(e))
    for name,key in (("source","sources"),("evidence","evidence"),("search-log","searches")):
        template=name+"-ledger-template.csv" if name!="search-log" else "search-log-template.csv"
        compare(f"ledgers/{name+'-ledger' if name!='search-log' else name}.csv",expected[key],fields(template))
    import openpyxl
    try:
        w=openpyxl.load_workbook(data/"网络检索与量化编码_Part04.03.xlsx",read_only=True,data_only=False)
        heads=template_headers()
        if w.sheetnames[:7]!=list(heads): errors.append("template sheet order mismatch")
        for name,header in heads.items():
            if [c.value for c in w[name][1]]!=header: errors.append("template header mismatch:"+name)
        for s in w:
            for row in s:
                if any(c.data_type=="f" for c in row): errors.append("unexpected formula:"+s.title)
        for name in ("七维评分","主题编码"):
            if any(any(v is not None for v in row) for row in w[name].iter_rows(min_row=2,values_only=True)): errors.append("analysis must remain empty:"+name)
        # Compare native comment text in the actual reopened workbook, including overflow.
        long={}
        for s in w:
            if s.title.startswith("长文本续表"):
                for rid,n,text in s.iter_rows(min_row=2,values_only=True): long.setdefault(rid,[]).append((n,text))
        fulltext={k:"".join(x for _,x in sorted(v)) for k,v in long.items()}
        for name,header,rows in tables:
            if name not in w:errors.append('missing entity sheet:'+name);continue
            values=list(w[name].values);reference=[header,*rows]
            if len(values)!=len(reference):errors.append('entity sheet row mismatch:'+name);continue
            for actual,ref in zip(values,reference):
                if len(actual)!=len(ref):errors.append('entity sheet column mismatch:'+name);continue
                for av,bv in zip(actual,ref):
                    av='' if av is None else av;bv='' if bv is None else bv
                    if isinstance(bv,str) and (xml_unsupported(bv) or len(bv.encode('utf-16-le'))//2>32767):
                        if isinstance(av,str) and '：' in av:
                            rid=av.split('：',1)[1].split('；',1)[0];saved=fulltext.get(rid)
                            av=json.loads(saved) if xml_unsupported(bv) and saved is not None else saved
                    if av!=bv:errors.append('entity sheet mismatch:'+name)
        loaded=[]
        for s in w:
            if s.title=="原始评论" or s.title.startswith("原始评论_"):
                h=[c.value for c in s[1]]
                for values in s.iter_rows(min_row=2,values_only=True):
                    record=dict(zip(h,values))
                    loaded.append(record)
        if len(loaded)!=len(expected["reviews"]): errors.append("workbook review count mismatch")
        else:
            for a,b in zip(loaded,expected["reviews"]):
                for key in REVIEW_FIELDS:
                    av,bv=a.get(key),b.get(key)
                    if isinstance(bv,bool): bv="是" if bv else "否"
                    if isinstance(bv,(dict,list)): bv=json.dumps(bv,ensure_ascii=False)
                    if av is None: av=""
                    if bv is None: bv=""
                    if isinstance(bv,str) and (xml_unsupported(bv) or len(bv.encode("utf-16-le"))//2>32767):
                        prefix="XML不可表示字符，JSON转义原文见" if xml_unsupported(bv) else "完整长文本见"
                        if isinstance(av,str) and av.startswith(prefix):
                            rid=av.split("：",1)[1].split("；",1)[0]
                            decoded=fulltext.get(rid)
                            av=json.loads(decoded) if xml_unsupported(bv) and decoded is not None else decoded
                    if av!=bv: errors.append("workbook raw field mismatch:"+key)
        w.close()
    except Exception as e: errors.append("workbook:"+str(e))
    finally:
        if w is not None: w.close()
    report={"status":"invalid" if errors else "valid","errors":errors,"reviews":len(expected["reviews"]),
            "sources":len(expected["sources"]),"checked_at":now(),"scope":"保存后完整性；不证明平台总体覆盖或正式评分资格"}
    write_json(data/"verify/完整性核验结果.json",report)
    return report
