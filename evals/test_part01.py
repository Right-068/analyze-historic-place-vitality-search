"""Anonymous offline contracts. Run: python -m unittest discover -s evals -v."""
import copy,csv,json,os,subprocess,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"scripts"))
import common,ctrip,spatial,poi_filter,poi_match,preflight,runner,export_data
from collect import Store,collect_site
from shapely.geometry import Polygon,Point,mapping
from pyproj import CRS
import shapefile

def candidate(pid="901",name="合成文化馆"):
    return {"poiId":pid,"sight_id":pid,"name":name,"aliases":"","city":"合成市","district":"","address":"测试街12号",
            "lat":23.01,"lon":113.01,"url":f"https://you.ctrip.com/sight/test/{pid}.html","type":"sight","commentCount":3}
def obj():
    return {"object_id":"P1","name":"合成文化馆","aliases":"","city":"合成市","district":"","address":"测试街12号",
            "lon":113.01,"lat":23.01,"level":"poi","area_ids":["AREA-X"],"category":"A","type":"博物馆","parent_id":""}
def payload(items=None,total=3):
    result={"items":items if items is not None else [{"commentId":i,"content":text,"score":5,"userInfo":{"userId":"private"},
                    "publishTime":"/Date(1700000000000)/"} for i,text in enumerate(["A","=SUM(1,2)","原文\n保留"],1)]}
    if total is not None: result["totalCount"]=total
    return json.dumps({"result":result},ensure_ascii=False).encode()
class Response:
    def __init__(self,raw,status=200,url=ctrip.COMMENT_API):
        self.content=raw; self.status_code=status; self.url=url; self.headers={"Content-Type":"application/json"}
class FakeHttp:
    def __init__(self,raw=None,blocked=False):
        self.calls=[]; self.observer=None; self.raw=raw or payload(); self.blocked=blocked
    def request(self,method,url,**kw):
        self.calls.append((method,url,kw))
        if method=="GET":
            c=candidate()
            d={"id":int(c["sight_id"]),"poiId":int(c["poiId"]),"word":c["name"],"type":"sight","cityName":c["city"],"districtName":"","address":c["address"],"lat":c["lat"],"lon":c["lon"],"url":c["url"]}
            r=Response(json.dumps({"data":[d]},ensure_ascii=False).encode(),url=url)
        else: r=Response(self.raw,430 if self.blocked else 200,url)
        if self.observer: self.observer(r,common.now(),0)
        if self.blocked and method=="POST": raise common.AccessBlocked("HTTP 430")
        return r
class Contracts(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.data=Path(self.temp.name); self.cfg=copy.deepcopy(common.CONFIG)
        self.cfg["http"]["min_interval_seconds"]=1
        self.cfg['llm']['mode']='host'
        # Synthetic model adapter for historic transport/export tests, not a production fallback.
        # Production relationship behavior and pending/failure handling have separate Part04 tests.
        import llm_entity_match
        original=llm_entity_match.resolve
        def model(data,requests,cfg,read_only=False,offline=False):
            result=original(data,requests,cfg,read_only=read_only,offline=offline)
            if read_only:return result
            for r in requests:
                if r['request_id'] in result:continue
                a=r['research_entity'];b=r['platform_entity']
                relation=('represents_area' if b['name'].endswith(('文化旅游区','步行街')) else 'part_of') if a['entity_type']=='research_area' else 'same_as' if a['name']==b['name'] else 'unrelated'
                d={'request_id':r['request_id'],'relation':relation,'confidence':0.95,'reason':'匿名合成模型夹具，仅用于验证运输和导出合同。'}
                llm_entity_match.save_receipt(data,r,d,cfg,'synthetic-test-model')
            return original(data,requests,cfg,read_only=read_only,offline=offline)
        self.model_patch=patch('llm_entity_match.resolve',side_effect=model);self.model_patch.start();self.addCleanup(self.model_patch.stop)
    def tearDown(self): self.temp.cleanup()
    def seed(self):
        task={"task_run_id":"SYNTHETIC","input_sha256":preflight.identity(self.cfg)}
        common.write_json(self.data/"task.json",task)
        m={"status":"matched","object":obj(),"selected":candidate(),"candidates":[]}
        common.write_json(self.data/"logs/运行状态.json",{"matches":{"P1":m},"collections":{},"status":"running"})
    def test_matching_no_first_rank(self):
        bad=candidate("902"); bad["city"]="异名市"
        status,picked,_=poi_match.choose(obj(),[bad,candidate()],self.cfg["matching"])
        self.assertEqual(status,"matched"); self.assertEqual(picked["poiId"],"901")
    def test_matching_ambiguous(self):
        self.assertEqual(poi_match.choose(obj(),[candidate("901"),candidate("902")],self.cfg["matching"])[0],"ambiguous")
    def test_matching_missing_and_identity(self):
        for changes in ({"city":"","district":""},{"url":"https://you.ctrip.com/sight/test/5.html"},{"poiId":""},{"name":"相似合成文化馆"},{"lat":0,"lon":0,"address":""}):
            with self.subTest(changes=changes): self.assertFalse(poi_match.evaluate(obj(),{**candidate(),**changes},self.cfg["matching"])[0])
    def test_matching_far_coordinates(self):
        self.assertFalse(poi_match.evaluate(obj(),{**candidate(),"lon":115},self.cfg["matching"])[0])
    def test_manual_mapping_cannot_force_conflict(self):
        cfg=copy.deepcopy(self.cfg["matching"]); cfg["manual_mappings"]={"P1":"902"}
        self.assertEqual(poi_match.choose(obj(),[candidate()],cfg)[0],"conflict")
    def test_many_to_one_duplicate_candidate(self):
        self.assertEqual(poi_match.choose(obj(),[candidate(),candidate()],self.cfg["matching"])[0],"matched")
    def test_response_rejects_invalid(self):
        for raw in (b"{}",b"null",b"<html>error</html>",b'{"result":{"totalCount":0}}',b'{"result":{"items":[{}],"totalCount":1}}',b'{"success":false,"result":{"items":[]}}'):
            with self.subTest(raw=raw): self.assertRaises(common.InvalidResponse,ctrip.result_of,raw)
    def test_missing_total_not_zero(self):
        self.assertIsNone(ctrip.items_of(payload([],None))[1])
    def test_business_block(self):
        self.assertRaises(common.AccessBlocked,ctrip.decode,b'{"message":"captcha"}')
    def test_native_text_translation_and_author(self):
        r=ctrip.parse_item({"commentId":1,"content":"平台译文","originContent":"original text","userInfo":{"userNick":"private"}},"901")
        self.assertEqual(r["用户原文"],"original text"); self.assertEqual(r["评论全文"],"平台译文")
        self.assertNotIn("private",r["作者ID"])
    def test_exact_raw_storage_and_resume(self):
        h=FakeHttp(); s=Store(self.data,self.cfg,h); first=collect_site(s,candidate())
        calls=len(h.calls); again=collect_site(Store(self.data,self.cfg,h),candidate())
        self.assertEqual(first["unique_collected"],3); self.assertEqual(again["status"],"completed"); self.assertEqual(len(h.calls),calls)
        self.assertEqual(next((self.data/"raw").rglob("*.bin")).read_bytes(),h.raw)
    def test_raw_hash_damage(self):
        s=Store(self.data,self.cfg,FakeHttp()); s.page(candidate(),0,3,1)
        next((self.data/"raw").rglob("*.bin")).write_bytes(b"bad")
        self.assertRaises(common.IntegrityError,s.page,candidate(),0,3,1)
    def test_block_stops_all_queries(self):
        h=FakeHttp(blocked=True); s=Store(self.data,self.cfg,h)
        self.assertRaises(common.AccessBlocked,s.page,candidate(),0,3,1)
        n=len(h.calls); self.assertRaises(common.AccessBlocked,s.search,obj()); self.assertEqual(len(h.calls),n)
    def test_empty_unknown_partial(self):
        h=FakeHttp(payload([],None)); result=collect_site(Store(self.data,self.cfg,h),candidate())
        self.assertEqual(result["status"],"partial"); self.assertIsNone(result["interface_reported_total"])
    def test_repeat_page_bounded(self):
        h=FakeHttp(payload(total=100)); self.cfg["ctrip"]["stars"]=[0]; self.cfg["ctrip"]["sort_orders"]=[3]
        result=collect_site(Store(self.data,self.cfg,h),candidate())
        self.assertEqual(result["status"],"partial"); self.assertEqual(len(h.calls),2)
    def test_zero_total_complete(self):
        h=FakeHttp(payload([],0)); r=collect_site(Store(self.data,self.cfg,h),candidate())
        self.assertEqual(r["status"],"interface_empty"); self.assertEqual(r["unique_collected"],0)
        self.assertEqual(len(h.calls),2)
    def test_full_export_reopen_no_formulas(self):
        self.seed(); collect_site(Store(self.data,self.cfg,FakeHttp()),candidate())
        export_data.export(self.data,self.cfg); v=export_data.verify(self.data,self.cfg)
        self.assertEqual(v["status"],"valid",v)
        rows=common.read_csv(self.data/"reviews/原始评论.csv")
        self.assertEqual([r["评论全文"] for r in rows],["A","=SUM(1,2)","原文\n保留"])
    def test_long_unicode_not_truncated(self):
        self.seed(); raw=payload([{"commentId":1,"content":"😀"*33000+"尾","score":5}],1)
        collect_site(Store(self.data,self.cfg,FakeHttp(raw)),candidate())
        export_data.export(self.data,self.cfg); self.assertEqual(export_data.verify(self.data,self.cfg)["status"],"valid")
    def test_export_tamper_detected(self):
        self.seed(); collect_site(Store(self.data,self.cfg,FakeHttp()),candidate()); export_data.export(self.data,self.cfg)
        p=self.data/"reviews/原始评论.csv"; rows=common.read_csv(p); rows[0]["评论全文"]="替换"
        common.write_rows(p,rows); self.assertEqual(export_data.verify(self.data,self.cfg)["status"],"invalid")
    def test_input_config_conflict(self):
        self.seed(); collect_site(Store(self.data,self.cfg,FakeHttp()),candidate())
        self.cfg["ctrip"]["source_type"]=2
        self.assertEqual(preflight.check(self.data,self.cfg)["status"],"passed")
        self.assertEqual(export_data.reconstruct(self.data,self.cfg)["reviews"],[])
        fresh=FakeHttp(); collect_site(Store(self.data,self.cfg,fresh),candidate()); self.assertEqual(len(fresh.calls),1)
    def test_atomic_write_does_not_destroy_existing(self):
        p=self.data/"data.json"; common.atomic_bytes(p,b"first")
        with patch("common.os.replace",side_effect=OSError("interrupted")): self.assertRaises(OSError,common.atomic_bytes,p,b"second")
        self.assertEqual(p.read_bytes(),b"first"); self.assertEqual(list(self.data.glob("*.tmp")),[])
    def test_task_lock(self):
        with common.task_lock(self.data): self.assertRaises(RuntimeError,common.task_lock(self.data).__enter__)
    def test_paths_confined(self): self.assertRaises(ValueError,common.task_path,self.data,"../outside")
    def test_classification_compound_parent(self):
        rules=common.read_json(common.SKILL_DIR/"assets/poi_filter_rules.json")
        for row,want in [({"name":"合成馆","typecode":"050100|140100","parent_id":"P"},"A"),
                         ({"name":"合成博物馆入口","typecode":"140100"},"C"),
                         ({"name":"合成老字号","typecode":"050100"},"B"),
                         ({"name":"普通住宅","typecode":"120100"},"D")]:
            with self.subTest(row=row): self.assertEqual(poi_filter.classify(row,"inside",rules)[0],want)
        self.assertEqual(poi_filter.classify({"name":"博物馆"},"boundary_uncertain",rules)[0],"E")
    def test_coordinate_roundtrip_and_wrong_assignment(self):
        x,y=spatial.gcj_forward(113.,23.); a,b=spatial.gcj_inverse(x,y)
        self.assertAlmostEqual(a,113.,places=7); self.assertAlmostEqual(b,23.,places=7)
        row={"longitude_gcj02":x,"latitude_gcj02":y,"longitude_cgcs2000_approx":x,"latitude_cgcs2000_approx":y}
        meter=spatial.metric_transform(Polygon([(113,23),(113.1,23),(113.1,23.1),(113,23.1)]))
        self.assertEqual(spatial.coordinate(row,meter,10)[1],"coordinate_conflict")
    def test_polygon_hole_boundary(self):
        p=Polygon([(0,0),(3,0),(3,3),(0,3)],holes=[[(1,1),(1,2),(2,2),(2,1)]])
        self.assertTrue(p.covers(Point(0,1))); self.assertFalse(p.covers(Point(1.5,1.5)))
    def test_real_scope_and_full_input(self):
        project,base,city,metre,scopes,zones,features=spatial.load_scope(self.cfg)
        self.assertEqual(len(scopes),23); self.assertEqual(sum(not z["selected"] for z in zones),3)
        self.assertEqual(len(common.read_csv(base/project["poi_file"])),47542)
        self.assertTrue(all(city.covers(s["geometry"]) for s in scopes.values()))
    def test_unattended_combined_runner_resume(self):
        fake=FakeHttp()
        with patch("runner.build",return_value=([obj()],{"districts":1,"search_objects":1})):
            r=runner.execute(self.data,self.cfg,fake); n=len(fake.calls)
            self.assertEqual(r["status"],"completed",r)
            r2=runner.execute(self.data,self.cfg,fake); self.assertEqual(len(fake.calls),n); self.assertEqual(r2["verification"],"valid")
    def test_unattended_blocked_finishes_without_prompt(self):
        with patch("runner.build",return_value=([obj()],{"districts":1})):
            r=runner.execute(self.data,self.cfg,FakeHttp(blocked=True))
            self.assertEqual(r["status"],"failed")
            self.assertTrue((self.data/"网络检索与量化编码_Part04.03.xlsx").exists())
    def test_cli_no_coding_entry(self):
        r=subprocess.run([sys.executable,"-B",str(common.SKILL_DIR/"scripts/hpv.py"),"--help"],capture_output=True,text=True,encoding="utf8")
        self.assertEqual(r.returncode,0)
        self.assertNotIn("validation-sample",r.stdout); self.assertNotIn("agreement",r.stdout)

    def test_name_normalization_dashes(self):
        self.assertEqual(spatial.normalize_name("合成—片区（甲）历史文化街区核心保护范围"),spatial.normalize_name("合成-片区(甲)历史文化街区建设控制地带"))
    def test_http_statuses_and_bounded_retries(self):
        for code,exc,calls in [(401,common.AccessBlocked,1),(403,common.AccessBlocked,1),(429,common.AccessBlocked,1),(430,common.AccessBlocked,1),(404,common.InvalidResponse,1),(500,common.InvalidResponse,3)]:
            with self.subTest(code=code):
                h=common.Http(self.cfg)
                with patch.object(h.s,"request",return_value=Response(b"{}",code)) as request,patch("common.time.sleep"):
                    self.assertRaises(exc,h.request,"GET",ctrip.SEARCH_URL.format("synthetic"))
                    self.assertEqual(request.call_count,calls)
    def test_transient_http_retry_then_success(self):
        h=common.Http(self.cfg)
        with patch.object(h.s,"request",side_effect=[Response(b"{}",500),Response(payload())]) as request,patch("common.time.sleep"):
            self.assertEqual(h.request("POST",ctrip.COMMENT_API).status_code,200)
            self.assertEqual(request.call_count,2)
    def test_capture_survives_interruption_before_commit(self):
        class Interrupted(FakeHttp):
            def request(self,*a,**kw):
                r=super().request(*a,**kw)
                raise KeyboardInterrupt()
        store=Store(self.data,self.cfg,Interrupted())
        self.assertRaises(KeyboardInterrupt,store.page,candidate(),0,3,1)
        rawfiles=list((self.data/"raw").rglob("*.bin")); self.assertEqual(len(rawfiles),1)
        before=rawfiles[0].read_bytes()
        result=collect_site(Store(self.data,self.cfg,FakeHttp()),candidate())
        self.assertEqual(result["status"],"completed"); self.assertEqual(rawfiles[0].read_bytes(),before)
    def test_incomplete_not_marked_complete(self):
        self.cfg["ctrip"]["max_pages_per_slice"]=1; self.cfg["ctrip"]["stars"]=[0]; self.cfg["ctrip"]["sort_orders"]=[3]
        r=collect_site(Store(self.data,self.cfg,FakeHttp(payload(total=10))),candidate())
        self.assertEqual(r["status"],"partial"); self.assertEqual(r["unavailable_count"],7)
    def test_partial_output_survives_one_failed_unit(self):
        self.seed()
        class Sometimes(FakeHttp):
            def request(self,method,url,**kw):
                if kw.get("json",{}).get("arg",{}).get("pageIndex")==2: raise common.InvalidResponse("network")
                return super().request(method,url,**kw)
        self.cfg["ctrip"]["stars"]=[0]; self.cfg["ctrip"]["sort_orders"]=[3]
        r=collect_site(Store(self.data,self.cfg,Sometimes(payload(total=10))),candidate())
        self.assertEqual(r["status"],"partial")
        export_data.export(self.data,self.cfg)
        self.assertEqual(export_data.verify(self.data,self.cfg)["reviews"],3)
    def test_duplicates_across_slices_preserve_occurrences(self):
        self.seed(); h=FakeHttp(payload(total=100))
        self.cfg["ctrip"]["stars"]=[0,1]; self.cfg["ctrip"]["sort_orders"]=[3]
        collect_site(Store(self.data,self.cfg,h),candidate())
        r=export_data.reconstruct(self.data,self.cfg)
        self.assertEqual(len(r["reviews"]),3); self.assertGreater(len(r["occurrences"]),3)
    def test_workbook_tamper_detected(self):
        self.seed(); collect_site(Store(self.data,self.cfg,FakeHttp()),candidate()); export_data.export(self.data,self.cfg)
        import openpyxl
        p=self.data/"网络检索与量化编码_Part04.03.xlsx"; w=openpyxl.load_workbook(p)
        h=[c.value for c in w["原始评论"][1]]; w["原始评论"].cell(2,h.index("评论全文")+1,"改写")
        w.save(p); w.close()
        self.assertEqual(export_data.verify(self.data,self.cfg)["status"],"invalid")
    def test_invalid_coordinate_nan(self):
        metre=spatial.metric_transform(Polygon([(113,23),(113.1,23),(113.1,23.1),(113,23.1)]))
        p,status,_=spatial.coordinate({"longitude_cgcs2000_approx":"nan","latitude_cgcs2000_approx":"23","longitude_gcj02":"113","latitude_gcj02":"23"},metre,10)
        self.assertEqual(status,"invalid_coordinate")
    def test_known_crs_matching(self):
        cfg=copy.deepcopy(self.cfg["matching"]); cfg["coordinate_crs"]="WGS84"
        o=obj(); o["address"]=""
        self.assertTrue(poi_match.evaluate(o,candidate(),cfg)[0])
    def test_raw_immutability_after_export(self):
        self.seed(); collect_site(Store(self.data,self.cfg,FakeHttp()),candidate())
        def hashes(): return {str(p):common.digest_file(p) for p in (self.data/"raw").rglob("*.bin")}
        before=hashes(); export_data.export(self.data,self.cfg); export_data.verify(self.data,self.cfg)
        self.assertEqual(hashes(),before)
    def test_empty_search_not_error(self):
        self.assertEqual(ctrip.search_candidates(b'{"data":[]}'),[])
    def test_platform_changed_text_is_retained(self):
        self.seed(); store=Store(self.data,self.cfg,FakeHttp())
        store.page(candidate(),0,3,1)
        store.http=FakeHttp(payload([{"commentId":1,"content":"later version","score":5}],1))
        store.page(candidate(),1,3,1)
        r=export_data.reconstruct(self.data,self.cfg)
        self.assertEqual(len(r["reviews"]),3); self.assertEqual(len(r["occurrences"]),4)
        self.assertFalse(r["errors"])

    def small_dataset(self):
        base=self.data/"input"; base.mkdir()
        from shapely.geometry.polygon import orient
        for name,bounds in (("city",(113,23,113.04,23.04)),("core",(113.002,23.002,113.018,23.018)),("control",(113.005,23.005,113.02,23.02))):
            from shapely.geometry import box
            with shapefile.Writer(str(base/name),shapeType=shapefile.POLYGON,encoding="utf8") as w:
                w.field("MC","C",120); w.shape(mapping(orient(box(*bounds),sign=-1))); w.record("合成区历史文化街区")
            (base/name).with_suffix(".prj").write_text(CRS.from_epsg(4490).to_wkt(),encoding="utf8")
        gx,gy=spatial.gcj_forward(113.01,23.01)
        row={"poi_id":"P1","name":"合成文化馆","type":"博物馆","typecode":"140100","longitude_gcj02":gx,"latitude_gcj02":gy,
             "longitude_cgcs2000_approx":113.01,"latitude_cgcs2000_approx":23.01,"parent_id":"","address":"测试街12号","district":"","alias":""}
        common.write_rows(base/"poi.csv",[row])
        common.write_json(base/"project.json",{"dataset_id":"synthetic","poi_file":"poi.csv","city_boundary":"city.shp","layers":{"core":"core.shp","control":"control.shp"},
            "excluded_districts":[],"expected_districts":1,"name_field":"MC","city":"合成市"})
        self.cfg["dataset"]=str(base/"project.json"); self.cfg["matching"]["coordinate_crs"]="WGS84"
        return base
    def test_actual_geometry_to_final_workbook_and_resume(self):
        self.small_dataset()
        from urllib.parse import parse_qs,urlparse
        class GeometricHttp(FakeHttp):
            def request(self,method,url,**kw):
                if method!="GET": return super().request(method,url,**kw)
                self.calls.append((method,url,kw)); name=parse_qs(urlparse(url).query)["keyword"][0]
                pid="901" if '文化馆' in name else "902"
                # Query variants return the same native identity, not renamed copies of one ID.
                name='合成文化馆' if pid=='901' else '合成区文化旅游区'
                c=candidate(pid,name)
                d={"id":int(pid),"poiId":int(pid),"word":name,"type":"sight","cityName":c["city"],"address":c["address"],"lat":c["lat"],"lon":c["lon"],"url":c["url"]}
                reply=Response(json.dumps({"data":[d]},ensure_ascii=False).encode(),url=url)
                if self.observer: self.observer(reply,common.now(),0)
                return reply
        fake=GeometricHttp(); task=self.data/"task"
        result=runner.execute(task,self.cfg,fake)
        self.assertEqual(result["status"],"completed",result); self.assertEqual(result["objects"],2); self.assertEqual(result["reviews"],6)
        levels={r["entity_level"] for r in common.read_csv(task/"ledgers/evidence-ledger.csv")}
        # A tourism expression is not proof that comments cover the entire legal area.
        self.assertEqual(levels,{"poi"})
        self.assertIn('represents_area',{m['relation_type'] for m in common.read_json(task/'logs/运行状态.json')['entity_mappings']})
        n=len(fake.calls); ids=[r["内部样本编号"] for r in common.read_csv(task/"reviews/原始评论.csv")]
        self.assertEqual(runner.execute(task,self.cfg,fake)["verification"],"valid"); self.assertEqual(len(fake.calls),n)
        self.assertEqual(ids,[r["内部样本编号"] for r in common.read_csv(task/"reviews/原始评论.csv")])
    def test_real_process_exit_and_new_process_state_recovery(self):
        r=subprocess.run([sys.executable,"-B",str(common.SKILL_DIR/"evals/interrupt_capture.py"),str(self.data)],capture_output=True)
        self.assertEqual(r.returncode,19,r.stderr)
        raw=list((self.data/"raw").rglob("*.bin")); before={str(p):common.digest_file(p) for p in raw}
        self.assertEqual(len(raw),1)
        units=common.read_json(self.data/"logs/采集单元.json"); self.assertEqual(next(iter(units.values()))["status"],"running")
        h=FakeHttp(); result=collect_site(Store(self.data,self.cfg,h),candidate())
        self.assertEqual(result["status"],"completed"); self.assertEqual(len(h.calls),0)
        self.assertTrue(all(common.digest_file(p)==sha for p,sha in before.items()))
        n=len(h.calls); collect_site(Store(self.data,self.cfg,h),candidate()); self.assertEqual(len(h.calls),n)
    def test_duplicate_categories_preserved_and_external_relation(self):
        base=self.small_dataset(); rows=common.read_csv(base/"poi.csv"); first=rows[0]; first.update(type="普通企业",typecode="170100")
        second={**first,"type":"博物馆","typecode":"140100"}
        gx,gy=spatial.gcj_forward(113.03,23.03)
        outside={**first,"poi_id":"P2","name":"合成附属设施","parent_id":"P1","longitude_gcj02":gx,"latitude_gcj02":gy,"longitude_cgcs2000_approx":113.03,"latitude_cgcs2000_approx":23.03}
        common.write_rows(base/"poi.csv",[first,second,outside])
        tasks,summary=spatial.build(self.data/"task",self.cfg)
        rows=common.read_csv(self.data/"task/scope/高德POI候选清单.csv")
        self.assertEqual(summary["poi_input_rows"],3); self.assertEqual(summary["unique_poi"],2)
        self.assertEqual(rows[0]["typecode"],"170100|140100"); self.assertEqual(rows[0]["candidate_class"],"A")
        self.assertEqual(rows[1]["object_level"],"external_related")
    def test_native_rating_workbook_tamper(self):
        self.seed(); collect_site(Store(self.data,self.cfg,FakeHttp()),candidate()); export_data.export(self.data,self.cfg)
        import openpyxl
        p=self.data/"网络检索与量化编码_Part04.03.xlsx"; w=openpyxl.load_workbook(p)
        h=[c.value for c in w["原始评论"][1]]; w["原始评论"].cell(2,h.index("总评分")+1,0); w.save(p); w.close()
        self.assertEqual(export_data.verify(self.data,self.cfg)["status"],"invalid")
    def test_matching_conflict_is_not_unmatched(self):
        self.assertEqual(poi_match.choose(obj(),[{**candidate(),"lon":115}],self.cfg["matching"])[0],"conflict")
    def test_configuration_rejects_wrong_types(self):
        for value in ([1],{"ctrip":{"page_size":True}},{"spatial":{"clip_to_city":"false"}},{"export":{"anonymize_author":False}}):
            p=self.data/"cfg.json"; common.write_json(p,value)
            with self.subTest(value=value): self.assertRaises(ValueError,preflight.config,p)
    def test_network_preflight_does_not_pollute_formal_units(self):
        self.small_dataset(); task=self.data/"task"; fake=FakeHttp()
        with patch("collect.Http",return_value=fake): result=preflight.check(task,self.cfg,network=True)
        self.assertEqual(result["status"],"passed",result)
        self.assertFalse((task/"logs/采集单元.json").exists())
        self.assertTrue((task/"verify/network-probe/logs/采集单元.json").exists())
    def test_preflight_restriction_stops_run_requests(self):
        self.small_dataset(); task=self.data/"task"; fake=FakeHttp(blocked=True)
        with patch("collect.Http",return_value=fake): result=preflight.check(task,self.cfg,network=True)
        self.assertEqual(result["status"],"blocked")
        n=len(fake.calls); result=runner.execute(task,self.cfg,fake)
        self.assertEqual(result["status"],"failed"); self.assertEqual(len(fake.calls),n)
    def test_zero_uncertainty_buffer_does_not_remove_inside_points(self):
        self.small_dataset(); self.cfg["spatial"]["boundary_uncertainty_m"]=0
        tasks,summary=spatial.build(self.data/"task",self.cfg)
        self.assertEqual(summary["search_objects"],2)
        self.assertEqual(common.read_csv(self.data/"task/scope/高德POI候选清单.csv")[0]["spatial_status"],"inside")
    def test_shapefile_multipart_and_hole_are_preserved(self):
        from shapely.geometry import MultiPolygon,box
        from shapely.geometry.polygon import orient
        p=self.data/"multi"
        polygon=Polygon([(113,23),(113.02,23),(113.02,23.02),(113,23.02)],holes=[[(113.005,23.005),(113.015,23.005),(113.015,23.015),(113.005,23.015)]])
        g=MultiPolygon([orient(polygon,sign=-1),orient(box(113.03,23.03,113.04,23.04),sign=-1)])
        with shapefile.Writer(str(p),shapeType=shapefile.POLYGON) as w:
            w.field("Id","N"); w.shape(mapping(g)); w.record(0)
        p.with_suffix(".prj").write_text(CRS.from_epsg(4490).to_wkt(),encoding="utf8")
        rows,_=spatial.read_layer(p.with_suffix(".shp")); parsed=rows[0][1]
        self.assertEqual(parsed.geom_type,"MultiPolygon"); self.assertFalse(parsed.covers(Point(113.01,23.01)))
        self.assertTrue(parsed.covers(Point(113.035,23.035)))
if __name__=="__main__": unittest.main()
